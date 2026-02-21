# import sys
# import io
# import fitz  # PyMuPDF
# from pptx import Presentation
# from pptx.util import Inches, Pt
# from pptx.dml.color import RGBColor


# def pdf_to_ppt(pdf_bytes):

#     doc = fitz.open(stream=pdf_bytes, filetype="pdf")
#     prs = Presentation()

#     for page in doc:

#         page_width = page.rect.width
#         page_height = page.rect.height

#         # ===== ORIENTATION =====
#         if page_height > page_width:
#             prs.slide_width = Inches(7.5)
#             prs.slide_height = Inches(10)
#         else:
#             prs.slide_width = Inches(13.33)
#             prs.slide_height = Inches(7.5)

#         slide = prs.slides.add_slide(prs.slide_layouts[6])

#         # ===== EXTRACT TEXT (SPAN LEVEL) =====
#         text_dict = page.get_text("dict")

#         for block in text_dict["blocks"]:

#             # ===== TEXT BLOCK =====
#             if block["type"] == 0:

#                 for line in block["lines"]:
#                     for span in line["spans"]:

#                         text = span["text"]
#                         if not text.strip():
#                             continue

#                         x0, y0, x1, y1 = span["bbox"]

#                         left = (x0 / page_width) * prs.slide_width
#                         top = (y0 / page_height) * prs.slide_height
#                         width = ((x1 - x0) / page_width) * prs.slide_width
#                         height = ((y1 - y0) / page_height) * prs.slide_height

#                         textbox = slide.shapes.add_textbox(
#                             left, top, width, height
#                         )

#                         tf = textbox.text_frame
#                         tf.clear()

#                         p = tf.paragraphs[0]
#                         p.text = text

#                         # ===== FONT SETTINGS =====
#                         font = p.font
#                         font.size = Pt(span["size"])

#                         if span.get("font"):
#                             font.name = span["font"]

#                         if "Bold" in span.get("font", ""):
#                             font.bold = True

#                         if "Italic" in span.get("font", ""):
#                             font.italic = True

#                         if span.get("color") is not None:
#                             color = span["color"]
#                             r = (color >> 16) & 255
#                             g = (color >> 8) & 255
#                             b = color & 255
#                             font.color.rgb = RGBColor(r, g, b)

#             # ===== IMAGE BLOCK =====
#             elif block["type"] == 1 and "image" in block:

#                 x0, y0, x1, y1 = block["bbox"]

#                 left = (x0 / page_width) * prs.slide_width
#                 top = (y0 / page_height) * prs.slide_height
#                 width = ((x1 - x0) / page_width) * prs.slide_width
#                 height = ((y1 - y0) / page_height) * prs.slide_height

#                 img_bytes = block["image"]
#                 img_stream = io.BytesIO(img_bytes)

#                 slide.shapes.add_picture(
#                     img_stream,
#                     left,
#                     top,
#                     width,
#                     height
#                 )

#     output_stream = io.BytesIO()
#     prs.save(output_stream)
#     output_stream.seek(0)

#     return output_stream.read()


# if __name__ == "__main__":
#     try:
#         pdf_input = sys.stdin.buffer.read()
#         ppt_output = pdf_to_ppt(pdf_input)
#         sys.stdout.buffer.write(ppt_output)
#     except Exception as e:
#         sys.stderr.write(str(e))
#         sys.exit(1)


import sys
import io
import fitz
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


def pdf_to_ppt(pdf_bytes):

    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    prs = Presentation()

    for page in doc:

        page_width = page.rect.width
        page_height = page.rect.height

        # ===== ORIENTATION =====
        if page_height > page_width:
            prs.slide_width = Inches(7.5)
            prs.slide_height = Inches(10)
        else:
            prs.slide_width = Inches(13.33)
            prs.slide_height = Inches(7.5)

        slide = prs.slides.add_slide(prs.slide_layouts[6])

        # ===== TEXT EXTRACTION =====
        text_dict = page.get_text("dict")

        for block in text_dict["blocks"]:

            if block["type"] == 0:

                for line in block["lines"]:

                    # Calculate full line bbox
                    x0 = min(span["bbox"][0] for span in line["spans"])
                    y0 = min(span["bbox"][1] for span in line["spans"])
                    x1 = max(span["bbox"][2] for span in line["spans"])
                    y1 = max(span["bbox"][3] for span in line["spans"])

                    left = (x0 / page_width) * prs.slide_width
                    top = (y0 / page_height) * prs.slide_height
                    width = ((x1 - x0) / page_width) * prs.slide_width
                    height = ((y1 - y0) / page_height) * prs.slide_height

                    textbox = slide.shapes.add_textbox(left, top, width, height)
                    tf = textbox.text_frame
                    tf.clear()

                    p = tf.paragraphs[0]

                    for span in line["spans"]:
                        run = p.add_run()
                        run.text = span["text"]

                        font = run.font
                        font.size = Pt(span["size"])

                        if span.get("font"):
                            font.name = span["font"]

                        if "Bold" in span.get("font", ""):
                            font.bold = True

                        if "Italic" in span.get("font", ""):
                            font.italic = True

                        if span.get("color") is not None:
                            color = span["color"]
                            r = (color >> 16) & 255
                            g = (color >> 8) & 255
                            b = color & 255
                            font.color.rgb = RGBColor(r, g, b)

        # ===== IMAGE EXTRACTION (BETTER METHOD) =====
        image_list = page.get_images(full=True)

        for img in image_list:

            xref = img[0]
            base_image = doc.extract_image(xref)
            image_bytes = base_image["image"]

            # Find image rectangle
            for rect in page.get_image_rects(xref):

                left = (rect.x0 / page_width) * prs.slide_width
                top = (rect.y0 / page_height) * prs.slide_height
                width = ((rect.x1 - rect.x0) / page_width) * prs.slide_width
                height = ((rect.y1 - rect.y0) / page_height) * prs.slide_height

                img_stream = io.BytesIO(image_bytes)

                slide.shapes.add_picture(
                    img_stream,
                    left,
                    top,
                    width,
                    height
                )

    output_stream = io.BytesIO()
    prs.save(output_stream)
    output_stream.seek(0)

    return output_stream.read()


if __name__ == "__main__":
    try:
        pdf_input = sys.stdin.buffer.read()
        ppt_output = pdf_to_ppt(pdf_input)
        sys.stdout.buffer.write(ppt_output)
    except Exception as e:
        sys.stderr.write(str(e))
        sys.exit(1)